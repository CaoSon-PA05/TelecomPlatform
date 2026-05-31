"""
Excel exporter using openpyxl.
Produces .xlsx files compatible with legacy CDR Analyzer import format.
Multi-sheet support matches the legacy exportAllData() output:
  Sheets: TTTB | LIST | Contact | IMEI | Location
"""

from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .base_exporter import BaseExporter

# Column header → field name mapping for each sheet
SHEET_COLUMNS: dict[str, list[tuple[str, str]]] = {
    "TTTB": [
        ("So dien thoai",   "phone_normalized"),
        ("Ho ten",          "full_name"),
        ("Ngay sinh",       "date_of_birth"),
        ("Dia chi",         "address"),
        ("So giay to",      "id_doc_number"),
        ("Ngay cap",        "id_issue_date"),
        ("Ngay kich hoat",  "activation_date"),
    ],
    "LIST": [
        ("TT",           "source_file_row"),
        ("So chu",       "owner_phone"),
        ("So lien he",   "contact_number"),
        ("Thoi gian",    "recorded_at"),
        ("Thoi luong",   "duration_seconds"),
        ("IMEI",         "imei"),
        ("Ma tinh",      "province_code_raw"),
        ("Loai",         "comm_type"),
        ("Dich vu",      "service_direction_raw"),
        ("Dia chi tram", "bts_address"),
        ("LAC",          "lac"),
        ("Cell",         "cell_id"),
    ],
    "Contact": [
        ("TT",          "index"),
        ("Phone",       "owner_phone"),
        ("So dien thoai", "contact_phone"),
        ("Tan suat",    "total_count"),
        ("Zalo",        "zalo_id"),
        ("Facebook",    "facebook_url"),
        ("Telegram",    "telegram_id"),
        ("Ghi chu",     "notes"),
    ],
    "IMEI": [
        ("TT",            "index"),
        ("Phone",         "owner_phone"),
        ("IMEI",          "imei"),
        ("Tan suat",      "interaction_count"),
        ("Model",         "device_model"),
        ("Thoi gian dung","usage_period"),
        ("Ghi chu",       "notes"),
    ],
    "Location": [
        ("TT",         "index"),
        ("LAC",        "lac"),
        ("Cell",       "cell_id"),
        ("Ma tinh",    "province_code_raw"),
        ("Ten tram",   "bts_address"),
        ("Tan suat",   "total_count"),
        ("Google Maps","google_maps_url"),
    ],
}

_HEADER_FILL = PatternFill("solid", fgColor="1A2233")
_HEADER_FONT = Font(bold=True, color="00F0FF", name="Calibri", size=10)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


class ExcelExporter(BaseExporter):

    def export(self, data: list[dict], filename: str, sheet_name: str = "Sheet1") -> BytesIO:
        """Single-sheet export — used for per-tab exports."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel sheet name max 31 chars

        if data:
            cols = [(k, k) for k in data[0].keys()]
            self._write_header_row(ws, cols)
            for row_idx, row in enumerate(data, start=2):
                for col_idx, (_, field) in enumerate(cols, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(field))
            self._auto_column_widths(ws)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def export_multi_sheet(self, sheets: dict[str, list[dict]]) -> BytesIO:
        """
        Multi-sheet export.
        `sheets` maps sheet name → list of row dicts.
        Column headers derived from SHEET_COLUMNS mapping, falls back to dict keys.
        """
        wb = openpyxl.Workbook()
        first = True

        for sheet_name, rows in sheets.items():
            if first:
                ws = wb.active
                ws.title = sheet_name[:31]
                first = False
            else:
                ws = wb.create_sheet(title=sheet_name[:31])

            cols = SHEET_COLUMNS.get(sheet_name)
            if not cols and rows:
                cols = [(k, k) for k in rows[0].keys()]
            if not cols:
                continue

            self._write_header_row(ws, cols)
            for row_idx, row in enumerate(rows, start=2):
                for col_idx, (_, field) in enumerate(cols, start=1):
                    ws.cell(row=row_idx, column=col_idx, value=row.get(field))
            self._auto_column_widths(ws)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return "xlsx"

    def _write_header_row(self, ws, columns: list[tuple[str, str]]) -> None:
        """Write styled header row to an openpyxl worksheet."""
        for col_idx, (header, _) in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN
        ws.row_dimensions[1].height = 20

    def _auto_column_widths(self, ws) -> None:
        """Set column widths based on content length."""
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    max_len = max(max_len, len(str(cell.value or "")))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 4, 52)
