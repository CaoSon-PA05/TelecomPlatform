"""
FLA (Financial Ledger Analyzer) Excel exporter.

Produces a 3-sheet workbook from filtered bank transaction records:
  Sheet 1 — "Giao Dich"   : full filtered transaction list
  Sheet 2 — "Thong Ke"    : summary statistics
  Sheet 3 — "Phan Tich Gio": hourly distribution (0–23h)
"""

from __future__ import annotations

import re
from collections import Counter
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, numbers

from .base_exporter import BaseExporter

_HEADER_FILL   = PatternFill("solid", fgColor="1A2233")
_HEADER_FONT   = Font(bold=True, color="00F0FF", name="Calibri", size=10)
_HEADER_ALIGN  = Alignment(horizontal="center", vertical="center")
_DANGER_FONT   = Font(bold=True, color="FF3B30", name="Calibri", size=10)
_AMOUNT_FORMAT = '#,##0'

_TXN_COLS    = ["Ngay Gio", "Ngan Hang", "So Tien (VND)", "Ma Lenh", "Noi Dung", "Canh Bao"]
_STATS_COLS  = ["Chi So", "Gia Tri"]
_HOURLY_COLS = ["Gio", "So Giao Dich", "Co Bat Thuong"]


def _parse_hour(date_str: str) -> int | None:
    m = re.search(r"\s(\d{2}):", date_str or "")
    return int(m.group(1)) if m else None


class FlaExporter(BaseExporter):

    def export(self, data: list[dict], filename: str = "FLA_Report.xlsx") -> BytesIO:
        """
        Build a 3-sheet FLA workbook from a list of transaction dicts.
        Expected dict keys: date, bank, amount, code, content, flag.
        """
        wb = openpyxl.Workbook()

        self._sheet_transactions(wb.active, data)
        self._sheet_stats(wb.create_sheet(), data)
        self._sheet_hourly(wb.create_sheet(), data)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    # ------------------------------------------------------------------

    def _sheet_transactions(self, ws, data: list[dict]) -> None:
        ws.title = "Giao Dich"
        self._write_header(ws, _TXN_COLS)

        for r, d in enumerate(data, start=2):
            ws.cell(r, 1, d.get("date", ""))
            ws.cell(r, 2, d.get("bank", ""))
            amt_cell = ws.cell(r, 3, d.get("amount", 0))
            amt_cell.number_format = _AMOUNT_FORMAT
            ws.cell(r, 4, d.get("code", ""))
            ws.cell(r, 5, d.get("content", ""))
            flag = d.get("flag", "N/A")
            flag_cell = ws.cell(r, 6, flag)
            if flag and flag != "N/A":
                flag_cell.font = _DANGER_FONT
            # Highlight large or anomalous amounts
            amount = d.get("amount", 0)
            if isinstance(amount, (int, float)) and amount >= 150_000_000:
                ws.cell(r, 3).font = _DANGER_FONT

        self._auto_widths(ws)

    def _sheet_stats(self, ws, data: list[dict]) -> None:
        ws.title = "Thong Ke"
        self._write_header(ws, _STATS_COLS)

        total       = len(data)
        anomalies   = sum(1 for d in data if d.get("flag", "N/A") != "N/A")
        total_amt   = sum(d.get("amount", 0) for d in data)
        large_txns  = sum(1 for d in data if (d.get("amount", 0) or 0) >= 150_000_000)
        bank_counts = Counter(d.get("bank", "") for d in data)

        rows: list[tuple[str, object]] = [
            ("Tong so giao dich",          total),
            ("Giao dich bat thuong",        anomalies),
            ("Giao dich gia tri lon (>150M)", large_txns),
            ("Tong gia tri (VND)",          total_amt),
        ]
        for bank, count in sorted(bank_counts.items(), key=lambda x: -x[1]):
            rows.append((f"So giao dich - {bank}", count))

        for r, (label, value) in enumerate(rows, start=2):
            ws.cell(r, 1, label)
            val_cell = ws.cell(r, 2, value)
            if label == "Tong gia tri (VND)" and isinstance(value, (int, float)):
                val_cell.number_format = _AMOUNT_FORMAT

        self._auto_widths(ws)

    def _sheet_hourly(self, ws, data: list[dict]) -> None:
        ws.title = "Phan Tich Gio"
        self._write_header(ws, _HOURLY_COLS)

        hour_counts = Counter[int]()
        hour_anomaly = Counter[int]()
        for d in data:
            hr = _parse_hour(d.get("date", ""))
            if hr is not None:
                hour_counts[hr] += 1
                if d.get("flag", "N/A") != "N/A":
                    hour_anomaly[hr] += 1

        DANGER_HOURS = {23, 0, 1, 2, 3, 4}
        for r, hr in enumerate(range(24), start=2):
            count = hour_counts.get(hr, 0)
            anom  = hour_anomaly.get(hr, 0)
            ws.cell(r, 1, f"{hr:02d}:00")
            ws.cell(r, 2, count)
            ws.cell(r, 3, anom)
            if hr in DANGER_HOURS and count > 0:
                for c in range(1, 4):
                    ws.cell(r, c).font = _DANGER_FONT

        self._auto_widths(ws)

    # ------------------------------------------------------------------

    @staticmethod
    def _write_header(ws, cols: list[str]) -> None:
        for c, header in enumerate(cols, start=1):
            cell = ws.cell(1, c, header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN
        ws.row_dimensions[1].height = 20

    @staticmethod
    def _auto_widths(ws) -> None:
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value or "")) for cell in col), default=8
            )
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 52)

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return "xlsx"
