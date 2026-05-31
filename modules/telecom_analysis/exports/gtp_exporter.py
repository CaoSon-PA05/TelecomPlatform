"""
GTP (Geospatial Telemetry Processor) Excel exporter.

Produces a 3-sheet workbook from filtered Cell-LAC telemetry records:
  Sheet 1 — "Lich Trinh Tram"  : full filtered record list
  Sheet 2 — "Bao Cao LAC"      : LAC zone frequency summary
  Sheet 3 — "Thong Ke Tin Hieu": signal strength distribution
"""

from __future__ import annotations

import re
from collections import Counter
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from .base_exporter import BaseExporter

_HEADER_FILL = PatternFill("solid", fgColor="1A2233")
_HEADER_FONT = Font(bold=True, color="00F0FF", name="Calibri", size=10)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
_WARN_FONT = Font(bold=True, color="FF3B30", name="Calibri", size=10)

_RECORDS_COLS = [
    "Thoi Gian", "LAC", "Cell ID", "Vi Tri Tram", "Nha Mang",
    "Tin Hieu", "Vi Do (Lat)", "Kinh Do (Lng)",
]
_LAC_COLS     = ["Vung LAC", "So Lan Xuat Hien", "Ti Le (%)"]
_SIGNAL_COLS  = ["Phan Loai", "Nguong (dBm)", "So Ban Ghi", "Ti Le (%)"]


def _dbm_from_strength(strength: str) -> int | None:
    m = re.search(r"-(\d+)", strength or "")
    return -int(m.group(1)) if m else None


class GtpExporter(BaseExporter):

    def export(self, data: list[dict], filename: str = "GTP_Report.xlsx") -> BytesIO:
        """
        Build a 3-sheet GTP workbook from a list of cell-record dicts.
        Expected dict keys: time, lac, cell, label, ip, strength, lat, lng.
        """
        wb = openpyxl.Workbook()

        self._sheet_records(wb.active, data)
        self._sheet_lac_summary(wb.create_sheet(), data)
        self._sheet_signal_stats(wb.create_sheet(), data)

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    # ------------------------------------------------------------------

    def _sheet_records(self, ws, data: list[dict]) -> None:
        ws.title = "Lich Trinh Tram"
        self._write_header(ws, _RECORDS_COLS)

        for r, d in enumerate(data, start=2):
            ws.cell(r, 1, d.get("time", ""))
            ws.cell(r, 2, d.get("lac", ""))
            ws.cell(r, 3, d.get("cell", ""))
            ws.cell(r, 4, d.get("label", ""))
            ws.cell(r, 5, d.get("ip", ""))
            ws.cell(r, 6, d.get("strength", ""))
            ws.cell(r, 7, d.get("lat"))
            ws.cell(r, 8, d.get("lng"))

        self._auto_widths(ws)

    def _sheet_lac_summary(self, ws, data: list[dict]) -> None:
        ws.title = "Bao Cao LAC"
        self._write_header(ws, _LAC_COLS)

        lac_counts = Counter(d.get("lac", "") for d in data)
        total = len(data) or 1
        for r, (lac, count) in enumerate(
            sorted(lac_counts.items(), key=lambda x: -x[1]), start=2
        ):
            ws.cell(r, 1, lac)
            ws.cell(r, 2, count)
            ws.cell(r, 3, f"{count / total * 100:.1f}%")

        self._auto_widths(ws)

    def _sheet_signal_stats(self, ws, data: list[dict]) -> None:
        ws.title = "Thong Ke Tin Hieu"
        self._write_header(ws, _SIGNAL_COLS)

        good = medium = weak = unknown = 0
        for d in data:
            dbm = _dbm_from_strength(d.get("strength", ""))
            if dbm is None:
                unknown += 1
            elif dbm >= -75:
                good += 1
            elif dbm >= -90:
                medium += 1
            else:
                weak += 1

        total = len(data) or 1
        rows = [
            ("Tot",         ">= -75",   good,    good / total),
            ("Trung binh",  "-75 ~ -90", medium, medium / total),
            ("Yeu",         "< -90",    weak,    weak / total),
            ("Khong xac dinh", "N/A",   unknown, unknown / total),
        ]
        for r, (label, threshold, count, pct) in enumerate(rows, start=2):
            ws.cell(r, 1, label)
            ws.cell(r, 2, threshold)
            ws.cell(r, 3, count)
            ws.cell(r, 4, f"{pct * 100:.1f}%")
            if label == "Yeu":
                for c in range(1, 5):
                    ws.cell(r, c).font = _WARN_FONT

        self._auto_widths(ws)

    # ------------------------------------------------------------------

    @staticmethod
    def _write_header(ws, cols: list[str]) -> None:
        for c, header in enumerate(cols, start=1):
            cell = ws.cell(1, c, header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _HEADER_ALIGN
        ws.row_dimensions[1].height = 20

    @staticmethod
    def _auto_widths(ws) -> None:
        for col in ws.columns:
            max_len = max(
                (len(str(cell.value or "")) for cell in col), default=8
            )
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 52)

    def get_content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def get_file_extension(self) -> str:
        return "xlsx"
