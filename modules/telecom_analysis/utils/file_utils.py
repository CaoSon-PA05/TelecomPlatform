"""File handling utilities — Excel reading, filename parsing."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

import openpyxl

from .constants import CDR_SHEET_NAME

_PHONE_10_RE = re.compile(r"(?<!\d)(0[3-9]\d{8})(?!\d)")
_PHONE_9_RE  = re.compile(r"(?<!\d)([3-9]\d{8})(?!\d)")
_EXPORT_ALL_RE = re.compile(r"(export_all|xuat_tat_ca)", re.I)


def read_excel_to_array(path: Path, sheet_name: Optional[str] = None) -> list[list]:
    """
    Load an Excel file and return the target sheet as a 2D list (0-indexed).
    Uses openpyxl in read-only + data-only mode for performance.

    Sheet selection priority:
      1. Explicit sheet_name argument
      2. CDR_SHEET_NAME constant ('Sheet2') — confirmed in real files
      3. First sheet in the workbook
    """
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)

    if sheet_name and sheet_name in wb.sheetnames:
        target = sheet_name
    elif CDR_SHEET_NAME in wb.sheetnames:
        target = CDR_SHEET_NAME
    else:
        target = wb.sheetnames[0]

    ws = wb[target]
    data = [list(row) for row in ws.iter_rows(values_only=True)]
    wb.close()
    return data


def extract_phone_from_filename(filename: str) -> Optional[str]:
    """
    Extract a Vietnamese phone number from a CDR filename.

    Supported patterns:
      '0_0382733506.xlsx'  → '0382733506'
      '1_0969619929.xlsx'  → '0969619929'
      'cdr_0382733506_2025.xlsx' → '0382733506'

    Returns normalized phone with leading '0', or None.
    """
    stem = Path(filename).stem

    # Try 10-digit with leading 0
    m = _PHONE_10_RE.search(stem)
    if m:
        return m.group(1)

    # Try 9-digit → prepend 0
    m = _PHONE_9_RE.search(stem)
    if m:
        candidate = "0" + m.group(1)
        if _PHONE_10_RE.match(candidate):
            return candidate

    return None


def is_export_all_file(filename: str) -> bool:
    """
    Return True when the filename looks like a previously exported
    'export_all' multi-sheet workbook (these require a different import path).
    """
    return bool(_EXPORT_ALL_RE.search(Path(filename).stem))


def get_file_size_kb(path: Path) -> float:
    """Return file size in kilobytes, rounded to 1 decimal."""
    return round(path.stat().st_size / 1024, 1)
